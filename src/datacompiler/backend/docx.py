"""docx.py — Rendu Word (.docx), EXCLUSIVEMENT depuis page.resolved.
Structure calquée sur markdown.py/html.py (un titre par page, un paragraphe
par ligne regroupée), mais les tableaux deviennent de vrais tableaux Word
plutôt que du texte tabulaire.

Nouvelle dépendance : python-docx (à ajouter à requirements.txt)."""

from pathlib import Path
from docx import Document as DocxDocument

from document import Document
from .layout import _grouper_mots_en_lignes


def render_docx(doc: Document, output_path: str) -> str:
    docx_doc = DocxDocument()

    for page in doc.pages:
        docx_doc.add_heading(f"Page {page.number}", level=2)

        mots = page.resolved.words
        if mots:
            for ligne in _grouper_mots_en_lignes(mots):
                texte = " ".join(m.output_text for m in ligne)
                if texte.strip():
                    docx_doc.add_paragraph(texte)
        else:
            docx_doc.add_paragraph().add_run("(page vide)").italic = True

        for table in page.resolved.tables:
            _ajouter_tableau(docx_doc, table.rows)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    docx_doc.save(output_path)
    return output_path


def _ajouter_tableau(docx_doc, rows):
    if not rows:
        return
    nb_colonnes = max(len(row) for row in rows)
    table = docx_doc.add_table(rows=0, cols=nb_colonnes)
    table.style = "Table Grid"

    for i, row in enumerate(rows):
        cellules = table.add_row().cells
        for j in range(nb_colonnes):
            valeur = row[j] if j < len(row) and row[j] is not None else ""
            cellules[j].text = str(valeur)
        if i == 0:
            for cellule in cellules:
                for paragraphe in cellule.paragraphs:
                    for run in paragraphe.runs:
                        run.bold = True
